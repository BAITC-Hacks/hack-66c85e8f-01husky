import io
import json
import shutil
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from types import SimpleNamespace

import pytest
from docx import Document

from app.config import get_settings
from app.models import Meeting, Task
from app.models.enums import MeetingSource, MeetingStatus, Urgency
from app.services.sed import MockSED


@pytest.fixture
def meeting(db, seed):
    meeting = Meeting(
        title="Бюджет",
        meeting_date=date(2026, 9, 23),
        source=MeetingSource.upload,
        status=MeetingStatus.confirmed,
        created_by=seed["admin"].id,
        summary="## Решение\n- Проверить смету",
    )
    db.add(meeting)
    db.flush()
    db.add(
        Task(
            meeting_id=meeting.id,
            text="Проверить смету",
            assignee_name="Юротдел",
            urgency=Urgency.high,
        )
    )
    db.commit()
    db.refresh(meeting)
    return meeting


def test_export_docx_authorization_and_cleanup(admin_client, user_client, client, meeting):
    url = f"/api/v1/meetings/{meeting.id}/export"
    assert client.get(url).status_code == 401
    assert user_client.get(url).status_code == 403
    response = admin_client.get(url, params={"lang": "kk"})
    assert response.status_code == 200
    doc = Document(io.BytesIO(response.content))
    assert doc.paragraphs[0].text == "Жиналыс хаттамасы"
    assert "Жоғары" in " ".join(c.text for row in doc.tables[0].rows for c in row.cells)
    assert list((get_settings().data_dir / "exports").iterdir()) == []
    assert admin_client.get(url, params={"format": "html"}).status_code == 422
    assert admin_client.get(url, params={"lang": "en"}).status_code == 422


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice required for PDF integration")
def test_export_pdf_and_not_ready(admin_client, meeting, db):
    response = admin_client.get(f"/api/v1/meetings/{meeting.id}/export?format=pdf")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")
    assert response.headers["content-type"] == "application/pdf"
    meeting.status = MeetingStatus.processing
    db.commit()
    assert admin_client.get(f"/api/v1/meetings/{meeting.id}/export").status_code == 409


def test_sed_permissions_idempotency_and_task_reference(
    admin_client, user_client, meeting, db, monkeypatch
):
    def convert(path):
        pdf = path.with_suffix(".pdf")
        pdf.write_bytes(b"%PDF-1.4\nfixture")
        return pdf

    monkeypatch.setattr("app.routers.exports.to_pdf", convert)
    url = f"/api/v1/meetings/{meeting.id}/sed"
    assert user_client.post(url).status_code == 403
    meeting.status = MeetingStatus.draft
    db.commit()
    assert admin_client.post(url).status_code == 409
    meeting.status = MeetingStatus.confirmed
    db.commit()
    response = admin_client.post(url)
    assert response.status_code == 200
    assert admin_client.post(url).json() == response.json()
    reference = response.json()["sed_ref"]
    db.refresh(meeting)
    assert meeting.sed_ref == reference
    assert all(task.sed_ref == reference for task in meeting.tasks)
    directory = get_settings().outbox_dir / str(meeting.id)
    metadata = json.loads((directory / "meta.json").read_text())
    assert metadata["sed_ref"] == reference
    assert "audio_path" not in metadata
    assert (directory / "protocol.pdf").read_bytes().startswith(b"%PDF-")


def test_mock_sed_concurrent_sequence(tmp_path):
    pdf = tmp_path / "source.pdf"
    pdf.write_bytes(b"%PDF-1.4\nfixture")
    adapter = MockSED(tmp_path / "outbox")

    def push(number):
        meeting = SimpleNamespace(
            id=number, title="Meeting", meeting_date=date(2026, 9, 23), output_language="ru"
        )
        return adapter.push_protocol(meeting, pdf)

    with ThreadPoolExecutor(max_workers=4) as pool:
        references = list(pool.map(push, [1, 2, 3, 4, 1]))
    assert len(set(references)) == 4
    assert references[0] == references[-1]


def test_failed_pdf_conversion_cleans_export_directory(admin_client, meeting, monkeypatch):
    def fail_conversion(path):
        raise RuntimeError("soffice failed")

    monkeypatch.setattr("app.routers.exports.to_pdf", fail_conversion)
    response = admin_client.get(f"/api/v1/meetings/{meeting.id}/export?format=pdf")
    assert response.status_code == 503
    assert list((get_settings().data_dir / "exports").iterdir()) == []


def test_failed_sed_delivery_cleans_export_and_preserves_db(admin_client, meeting, monkeypatch, db):
    def convert(path):
        pdf = path.with_suffix(".pdf")
        pdf.write_bytes(b"%PDF-1.4\nfixture")
        return pdf

    def fail_delivery(*args):
        raise OSError("outbox unavailable")

    monkeypatch.setattr("app.routers.exports.to_pdf", convert)
    monkeypatch.setattr("app.routers.exports.MockSED.push_protocol", fail_delivery)
    with pytest.raises(OSError, match="outbox unavailable"):
        admin_client.post(f"/api/v1/meetings/{meeting.id}/sed")
    db.rollback()
    db.refresh(meeting)
    assert meeting.sed_ref is None
    assert all(task.sed_ref is None for task in meeting.tasks)
    assert list((get_settings().data_dir / "exports").iterdir()) == []
