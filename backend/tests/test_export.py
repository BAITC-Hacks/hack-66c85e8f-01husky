from datetime import date
from types import SimpleNamespace

import pytest
from docx import Document
from pipeline.fake import process
from pipeline.models import Participant

from app.services.export import build_docx, to_pdf


@pytest.fixture
def detail(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "app.services.export.get_settings", lambda: SimpleNamespace(data_dir=tmp_path)
    )
    participants = [
        Participant(id=1, name="Серик"),
        Participant(id=2, name="Дана"),
        Participant(id=3, name="Айбек"),
    ]
    result = process("sample.wav", date(2026, 9, 23), participants, ["Финансы", "ИТ"])
    return {
        "meeting": {"id": 42, "title": "Планирование", "meeting_date": date(2026, 9, 23)},
        "participants": participants,
        "summary": result.summary,
        "tasks": result.tasks,
        "segments": result.segments,
        "speaker_map": result.speaker_map,
    }


@pytest.mark.parametrize(
    ("lang", "title", "summary_heading", "urgency"),
    [
        ("ru", "Протокол совещания", "Краткое содержание", "Высокая"),
        ("kk", "Жиналыс хаттамасы", "Қысқаша мазмұны", "Жоғары"),
    ],
)
def test_build_docx_from_fake_meeting_detail(detail, lang, title, summary_heading, urgency):
    path = build_docx(detail, lang)
    doc = Document(path)
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    table_text = " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert title in paragraphs[0]
    assert "Планирование" in paragraphs[1]
    assert summary_heading in paragraphs
    assert "Подготовить обновлённую смету" in table_text
    assert "Айбек" in table_text
    assert urgency in table_text
    assert len(doc.tables[0].rows) == 5
    assert any("[00:00] Серик:" in p for p in paragraphs)
    assert not any(p.startswith("## ") for p in paragraphs)
    assert build_docx(detail, lang) != path


def test_export_rejects_unsupported_language(detail):
    with pytest.raises(ValueError, match="lang"):
        build_docx(detail, "en")


def test_to_pdf_requires_libreoffice(tmp_path, monkeypatch):
    docx_path = tmp_path / "protocol.docx"
    Document().save(docx_path)
    monkeypatch.setattr("app.services.export.shutil.which", lambda _: None)
    with pytest.raises(RuntimeError, match="soffice"):
        to_pdf(docx_path)


def test_to_pdf_does_not_accept_stale_output(tmp_path, monkeypatch):
    docx_path = tmp_path / "protocol.docx"
    Document().save(docx_path)
    docx_path.with_suffix(".pdf").write_bytes(b"stale")
    monkeypatch.setattr("app.services.export.shutil.which", lambda _: "/bin/soffice")
    monkeypatch.setattr("app.services.export.subprocess.run", lambda *args, **kwargs: None)
    with pytest.raises(RuntimeError, match="without creating"):
        to_pdf(docx_path)
