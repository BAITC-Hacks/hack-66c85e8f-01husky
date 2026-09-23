"""DOCX and PDF export for meeting protocols."""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from collections.abc import Mapping
from datetime import date, datetime
from enum import Enum
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor

from app.config import get_settings
from app.i18n import labels


def _value(item: Any, *keys: str, default: Any = None) -> Any:
    """Read a value from a dict or a pydantic/ORM object."""
    for key in keys:
        if isinstance(item, Mapping) and key in item:
            return item[key]
        if hasattr(item, key):
            return getattr(item, key)
    return default


def _display(value: Any, fallback: str = "—") -> str:
    if value is None or value == "":
        return fallback
    if isinstance(value, (date, datetime)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, Enum):
        return str(value.value)
    # XML 1.0 rejects these controls; an edited title or transcript may contain them.
    return re.sub(r"[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]", "", str(value))


def _as_list(value: Any) -> list[Any]:
    return list(value) if value is not None else []


def _task_direction(task: Any) -> str:
    direction = _value(task, "direction", "direction_name")
    if direction is None:
        direction = _value(_value(task, "direction_obj"), "name")
    return _display(_value(direction, "name", default=direction))


def _participant_label(person: Any, include_role: bool = True) -> str:
    name = _display(_value(person, "name"))
    role = _value(person, "position", "role")
    return f"{name} ({_display(role)})" if include_role and role else name


def _add_summary(document: Any, summary: str) -> None:
    for line in summary.splitlines():
        line = line.strip()
        if not line:
            continue
        heading = re.match(r"^#{1,6}\s+(.+)", line)
        if heading:
            paragraph = document.add_heading(level=2)
            line = heading[1]
        elif line.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="List Bullet")
            line = line[2:]
        else:
            paragraph = document.add_paragraph()
        for part in re.split(r"(\*\*[^*]+\*\*)", line):
            bold = part.startswith("**") and part.endswith("**")
            paragraph.add_run(part[2:-2] if bold else part).bold = bold


def build_docx(meeting_detail: Any, lang: str) -> Path:
    """Build a formatted protocol and return its path."""
    text = labels(lang)
    meeting = _value(meeting_detail, "meeting", default=meeting_detail)
    meeting_id = _value(meeting, "id", default="meeting")
    safe_id = re.sub(r"[^A-Za-z0-9_-]", "_", str(meeting_id))
    settings = get_settings()

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    for style in ("Title", "Heading 1", "Heading 2"):
        document.styles[style].font.color.rgb = RGBColor(0, 0, 0)
        for border in document.styles[style].element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)

    organization = document.add_paragraph()
    organization.alignment = WD_ALIGN_PARAGRAPH.CENTER
    organization_run = organization.add_run(text["organization"])
    organization_run.bold = True
    organization_run.font.size = Pt(11)
    if organization_name := getattr(settings, "organization_name", ""):
        organization_line = document.add_paragraph(_display(organization_name))
        organization_line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading = document.add_paragraph(style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = _display(_value(meeting, "title"), text["organization"])
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph(f"{text['date']}: {_display(_value(meeting, 'meeting_date', 'date'))}")

    participants = _as_list(_value(meeting_detail, "participants", default=[]))
    participant_names = [
        _participant_label(person) for person in participants if _value(person, "name")
    ]
    document.add_paragraph(
        f"{text['participants']}: {'; '.join(participant_names) if participant_names else '—'}"
    )

    summary = _value(meeting_detail, "summary", default=_value(meeting, "summary"))
    document.add_heading(text["summary"], level=1)
    _add_summary(document, _display(summary, ""))

    document.add_heading(text["tasks"], level=1)
    task_rows = _as_list(_value(meeting_detail, "tasks", default=[]))
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.autofit = False
    widths = (1.1, 3.0, 5.2, 2.1, 2.6, 3.1)
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Cm(width)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    headers = ["number", "assignee", "task", "deadline", "urgency", "direction"]
    for cell, key in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text[key]
        for paragraph in cell.paragraphs:
            for cell_run in paragraph.runs:
                cell_run.bold = True
    for index, task in enumerate(task_rows, 1):
        assignee = _value(task, "assignee_name")
        if not assignee:
            assignee = _value(task, "assignee", "participant_name")
        values = (
            str(index),
            _display(_value(assignee, "name", default=assignee)),
            _display(_value(task, "text", "title")),
            _display(_value(task, "deadline")),
            text.get(_display(_value(task, "urgency")), _display(_value(task, "urgency"))),
            _task_direction(task),
        )
        for cell, value in zip(table.add_row().cells, values, strict=True):
            cell.text = value
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = Cm(width)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for cell_run in paragraph.runs:
                    cell_run.font.size = Pt(9)

    segments = _as_list(_value(meeting_detail, "segments", default=[]))
    if segments:
        document.add_page_break()
        document.add_heading(text["transcript"], level=1)
        participant_by_id = {_value(person, "id"): person for person in participants}
        speaker_people = {
            str(_value(mapping, "speaker")): participant_by_id.get(
                _value(mapping, "participant_id")
            )
            for mapping in _as_list(_value(meeting_detail, "speaker_map", default=[]))
            if _value(mapping, "participant_id") is not None
        }
        seen_speakers = set()
        for segment in segments:
            start = float(_value(segment, "start", default=0))
            speaker = str(_value(segment, "speaker", default="Speaker"))
            person = speaker_people.get(speaker)
            label = (
                _participant_label(person, include_role=speaker not in seen_speakers)
                if person is not None
                else speaker
            )
            seen_speakers.add(speaker)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.add_run(
                f"[{int(start // 60):02d}:{int(start % 60):02d}] {label}: "
            ).bold = True
            paragraph.add_run(_display(_value(segment, "text"), ""))

    export_dir = settings.data_dir / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    output_dir = Path(tempfile.mkdtemp(prefix=f"{safe_id}-", dir=export_dir))
    output_path = output_dir / f"protocol-{lang}.docx"
    try:
        document.save(output_path)
    except Exception:
        shutil.rmtree(output_dir, ignore_errors=True)
        raise
    return output_path


def to_pdf(docx_path: str | Path) -> Path:
    """Convert a DOCX file to PDF with LibreOffice headless."""
    source = Path(docx_path).resolve()
    if not source.is_file() or source.suffix.lower() != ".docx":
        raise ValueError("docx_path must point to an existing .docx file")
    executable = shutil.which("soffice")
    if executable is None:
        raise RuntimeError("LibreOffice 'soffice' executable is required for PDF export")
    output = source.with_suffix(".pdf")
    output.unlink(missing_ok=True)
    with tempfile.TemporaryDirectory(prefix="protocol-lo-") as profile:
        subprocess.run(
            [
                executable,
                f"-env:UserInstallation={Path(profile).as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(source.parent),
                str(source),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=120,
        )
    if not output.is_file() or output.stat().st_size == 0:
        raise RuntimeError("LibreOffice completed without creating a PDF")
    with output.open("rb") as converted:
        if converted.read(5) != b"%PDF-":
            raise RuntimeError("LibreOffice output is not a PDF")
    return output
